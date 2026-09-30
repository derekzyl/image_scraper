"""Save receipt rows to an Excel file that keeps the leading 0."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from extractor import CSV_COLUMNS, ReceiptRow

HEADER = [title for _key, title in CSV_COLUMNS]
_RECIPIENT_COLUMN = 4  # 1-based column in the sheet


def save_rows(path: str | Path, rows: list[ReceiptRow], mode: str) -> dict[str, int | str]:
    """Write rows to an Excel workbook.

    mode "new" replaces the file. mode "append" adds rows and skips receipt
    numbers that are already in the file. The recipient number is stored as
    text so Excel and LibreOffice keep the leading 0.
    """
    destination = _workbook_path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    if mode not in {"new", "append"}:
        raise ValueError("Choose a new file or add to an existing one.")

    append = mode == "append" and destination.exists() and destination.stat().st_size > 0
    if append:
        workbook = load_workbook(destination)
        sheet = workbook.active
        _ensure_header(sheet)
        existing_ids = _existing_receipt_numbers(sheet)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADER)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        existing_ids = set()

    written = 0
    skipped = 0
    for row in rows:
        receipt_id = row.receipt_number.strip()
        if append and receipt_id and receipt_id in existing_ids:
            skipped += 1
            continue
        values = [row.as_csv()[key] for key, _title in CSV_COLUMNS]
        sheet.append(values)
        _keep_leading_zero(sheet, sheet.max_row)
        written += 1
        if receipt_id:
            existing_ids.add(receipt_id)

    _fit_columns(sheet)
    workbook.save(destination)
    return {"written": written, "skipped": skipped, "path": str(destination)}


def _workbook_path(path: str | Path) -> Path:
    destination = Path(path)
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")
    return destination


def _keep_leading_zero(sheet, row_index: int) -> None:
    cell = sheet.cell(row_index, _RECIPIENT_COLUMN)
    cell.number_format = "@"
    cell.value = "" if cell.value is None else str(cell.value)
    cell.alignment = Alignment(horizontal="left")


def _ensure_header(sheet) -> None:
    header = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    if header[: len(HEADER)] != HEADER:
        expected = ", ".join(HEADER)
        found = ", ".join(header)
        raise ValueError(
            "That file uses different columns, so rows were not added.\n"
            f"Expected: {expected}\n"
            f"Found: {found}"
        )


def _existing_receipt_numbers(sheet) -> set[str]:
    found: set[str] = set()
    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        receipt_id = "" if row[0] is None else str(row[0]).strip()
        if receipt_id:
            found.add(receipt_id)
    return found


def _fit_columns(sheet) -> None:
    widths = [22, 32, 28, 20, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].column_width = width
